import os
import subprocess
from pdf2docx import Converter
from pdf2image import convert_from_path
from pptx import Presentation
import camelot
from reportlab.pdfgen import canvas
from PIL import Image
from io import BytesIO
from reportlab.lib.pagesizes import letter
from pptx.util import Inches
from django.conf import settings
import pandas as pd

# Convert PDF to Word
def convert_pdf_to_word(pdf_file):
    pdf_path = os.path.join(settings.MEDIA_ROOT, pdf_file.name)

    # Save the uploaded PDF file to disk
    with open(pdf_path, 'wb+') as destination:
        for chunk in pdf_file.chunks():
            destination.write(chunk)

    # Define the output path for the .docx file
    docx_path = pdf_path.replace('.pdf', '.docx')

    # Convert the PDF to Word
    cv = Converter(pdf_path)
    cv.convert(docx_path, start=0, end=None)
    cv.close()

    return os.path.basename(docx_path)

# Convert Word to PDF (Linux-compatible using LibreOffice)
def convert_word_to_pdf(word_file):
    word_path = os.path.join(settings.MEDIA_ROOT, word_file.name)

    # Save the uploaded Word file to disk
    with open(word_path, 'wb+') as destination:
        for chunk in word_file.chunks():
            destination.write(chunk)

    # Use LibreOffice to convert Word to PDF
    try:
        subprocess.run([
            "libreoffice",
            "--headless",
            "--convert-to", "pdf",
            word_path,
            "--outdir", settings.MEDIA_ROOT
        ], check=True)
    except subprocess.CalledProcessError as e:
        raise RuntimeError("Conversion failed: " + str(e))

    # Get the output filename
    pdf_filename = word_file.name.replace('.docx', '.pdf')
    return pdf_filename



def convert_pdf_to_ppt(pdf_file):
    # Save PDF to disk
    pdf_path = os.path.join(settings.MEDIA_ROOT, pdf_file.name)
    with open(pdf_path, 'wb+') as destination:
        for chunk in pdf_file.chunks():
            destination.write(chunk)

    # Convert PDF pages to images
    images = convert_from_path(pdf_path)

    # Create PowerPoint
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]  # Blank layout

    for image in images:
        slide = prs.slides.add_slide(blank_slide_layout)
        img_path = pdf_path.replace('.pdf', '_temp.jpg')
        image.save(img_path, 'JPEG')
        slide.shapes.add_picture(img_path, Inches(0), Inches(0),
                                 width=prs.slide_width,
                                 height=prs.slide_height)

    pptx_filename = pdf_file.name.replace('.pdf', '.pptx')
    pptx_path = os.path.join(settings.MEDIA_ROOT, pptx_filename)
    prs.save(pptx_path)

    return os.path.basename(pptx_path)

def convert_ppt_to_pdf(ppt_file):
    ppt_path = os.path.join(settings.MEDIA_ROOT, ppt_file.name)

    # Save the uploaded PowerPoint file to disk
    with open(ppt_path, 'wb+') as destination:
        for chunk in ppt_file.chunks():
            destination.write(chunk)

    # Use LibreOffice to convert PPT to PDF
    try:
        subprocess.run([
            "libreoffice",
            "--headless",
            "--convert-to", "pdf",
            ppt_path,
            "--outdir", settings.MEDIA_ROOT
        ], check=True)
    except subprocess.CalledProcessError as e:
        raise RuntimeError("Conversion failed: " + str(e))

    # Get the output filename
    pdf_filename = ppt_file.name.replace('.pptx', '.pdf')
    return pdf_filename


# Convert PDF to Excel

def convert_pdf_to_excel(pdf_file):
    pdf_path = os.path.join(settings.MEDIA_ROOT, pdf_file.name)

    # Save the uploaded PDF file to disk
    with open(pdf_path, 'wb+') as destination:
        for chunk in pdf_file.chunks():
            destination.write(chunk)

    # Try extracting tables from the PDF using camelot
    tables = camelot.read_pdf(pdf_path, pages='all', flavor='stream')

    # If no tables are found, raise an error
    if len(tables) == 0:
        raise ValueError("No tables found in the PDF.")

    # Save the first table as an Excel file
    excel_filename = pdf_file.name.replace('.pdf', '.xlsx')
    excel_path = os.path.join(settings.MEDIA_ROOT, excel_filename)

    # Use pandas to save the first table in Excel format
    tables[0].df.to_excel(excel_path, index=False, engine='openpyxl')

    return os.path.basename(excel_path)


# Function to convert JPG to PDF
def convert_jpg_to_pdf(jpg_file):
    # Save the uploaded JPG file to disk
    image_path = os.path.join(settings.MEDIA_ROOT, jpg_file.name)
    with open(image_path, 'wb+') as destination:
        for chunk in jpg_file.chunks():
            destination.write(chunk)

    # Open the image using PIL
    image = Image.open(image_path)

    # Convert the image to RGB (remove alpha channel if it exists)
    if image.mode == 'RGBA':
        image = image.convert('RGB')

    # Define the output path for the PDF file
    pdf_filename = image_path.replace('.jpg', '.pdf')
    pdf_path = os.path.join(settings.MEDIA_ROOT, pdf_filename)

    # Save the image as a PDF
    image.save(pdf_path, 'PDF', resolution=100.0)

    return os.path.basename(pdf_path)